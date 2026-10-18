import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set standard page size A4 & Margins
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79) # 20mm
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98) # 25mm
    section.right_margin = Inches(0.79)

# Base Style Settings
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(13)
style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
style.paragraph_format.line_spacing = 1.3
style.paragraph_format.space_after = Pt(4)

with open("docs/BAO_CAO_THUC_HANH_TKQTPM.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

in_table = False
table_lines = []
in_code = False
code_lines = []

for line in lines:
    raw = line.rstrip('\r\n')
    
    # Handle Code Blocks
    if raw.startswith("```"):
        if in_code:
            # End code block
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("\n".join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
            # Add light border/bg via xml if needed, or keep clean
            code_lines = []
            in_code = False
        else:
            in_code = True
            code_lines = []
        continue
    
    if in_code:
        code_lines.append(raw)
        continue

    # Handle Markdown Tables
    if "|" in raw and ("---" in raw or in_table or raw.strip().startswith("|")):
        if "---" in raw:
            continue
        table_lines.append(raw)
        in_table = True
        continue
    else:
        if in_table:
            # Render Markdown Table
            if len(table_lines) > 0:
                rows_data = [[c.strip() for c in r.strip('|').split('|')] for r in table_lines]
                num_rows = len(rows_data)
                num_cols = max(len(r) for r in rows_data)
                tbl = doc.add_table(rows=num_rows, cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, row in enumerate(rows_data):
                    for c_idx in range(num_cols):
                        val = row[c_idx] if c_idx < len(row) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        cell.text = val
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.15
                        set_cell_margins(cell, 80, 80, 100, 100)
                        
                        if r_idx == 0:
                            set_cell_background(cell, "E2E8F0")
                            for r in p.runs:
                                r.font.bold = True
                                r.font.size = Pt(11)
                        else:
                            for r in p.runs:
                                r.font.size = Pt(11)
                
                # set borders for table
                tblPr = tbl._tbl.tblPr
                tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:insideV w:val="none"/></w:tblBorders>')
                tblPr.append(tblBorders)
                
            in_table = False
            table_lines = []

    if not raw.strip():
        continue

    # Title & Headings
    if raw.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(raw[2:].strip())
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    elif raw.startswith("## "):
        p = doc.add_paragraph()
        if "CHƯƠNG" in raw:
            p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(raw[3:].strip())
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
    elif raw.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(raw[4:].strip())
        run.bold = True
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    elif raw.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(raw[5:].strip())
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    # Image
    elif raw.startswith("!["):
        m = re.search(r'!\[(.*?)\]\((.*?)\)', raw)
        if m:
            caption = m.group(1)
            img_rel = m.group(2)
            img_path = os.path.join("docs", img_rel)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.add_run().add_picture(img_path, width=Inches(6.2))
    # Italic Caption
    elif raw.startswith("*Hình ") or raw.startswith("*"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(raw.strip('*_'))
        run.italic = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    # Lists
    elif raw.startswith("- ") or raw.startswith("* "):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        # Parse inline bold
        parts = re.split(r'(\*\*.*?\*\*)', raw[2:])
        for pt in parts:
            if pt.startswith("**") and pt.endswith("**"):
                r = p.add_run(pt[2:-2])
                r.bold = True
            else:
                p.add_run(pt)
    elif re.match(r'^\d+\.\s', raw):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        text_content = re.sub(r'^\d+\.\s*', '', raw)
        parts = re.split(r'(\*\*.*?\*\*)', text_content)
        for pt in parts:
            if pt.startswith("**") and pt.endswith("**"):
                r = p.add_run(pt[2:-2])
                r.bold = True
            else:
                p.add_run(pt)
    else:
        # Standard Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parts = re.split(r'(\*\*.*?\*\*)', raw)
        for pt in parts:
            if pt.startswith("**") and pt.endswith("**"):
                r = p.add_run(pt[2:-2])
                r.bold = True
            else:
                p.add_run(pt)

out_docx = "docs/BAO_CAO_THUC_HANH_TKQTPM.docx"
doc.save(out_docx)
print(f"Generated DOCX successfully: {out_docx}")
